"""
Generación de la factura/comprobante de un pedido en Excel (.xlsx), como
alternativa al PDF que ya genera app/utils/pdf.py — mismos datos y el mismo
desglose de Subtotal / IVA / Total, para quien prefiera abrirla o importarla
en una hoja de cálculo en vez de un PDF.

Se usa openpyxl (100% Python, sin dependencias de sistema) igual que
ReportLab para los PDF.
"""

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.pdf import (
    CORREO_EMPRESA,
    DIRECCION_EMPRESA,
    NIT_EMPRESA,
    NOMBRE_EMPRESA,
    REGIMEN_EMPRESA,
    TELEFONO_EMPRESA,
    formatear_fecha,
    formatear_moneda,
    lineas_totales_pedido,
)

NAVY = "0B1120"
CYAN = "0891B2"
SLATE = "475569"
GRIS_CLARO = "F1F5F9"
BLANCO = "FFFFFF"
TEXTO = "0F172A"


def generar_excel_comprobante(pedido: dict) -> bytes:
    # Este mismo generador se usa tanto para el comprobante de un pedido
    # como para una factura persistida (colección 'facturas'): una factura
    # trae 'pedidoId' apuntando al pedido del que se originó.
    id_pedido_relacionado = pedido.get("pedidoId") or pedido.get("id")
    numero_pedido = (id_pedido_relacionado or "")[-8:].upper() or "—"
    numero_factura = pedido.get("numero") or f"FE-{numero_pedido}"
    es_factura = pedido.get("pedidoId") is not None
    usuario = pedido.get("usuario") or {}
    nombre_cliente = (
        f"{usuario.get('nombre', '')} {usuario.get('apellido', '')}".strip() or "Usuario eliminado"
    )

    documento_cliente = "—"
    if usuario.get("numeroDocumento"):
        documento_cliente = f"{usuario.get('tipoDocumento') or ''} {usuario['numeroDocumento']}".strip()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Factura {numero_pedido}"[:31]

    fuente_titulo = Font(name="Calibri", size=16, bold=True, color=NAVY)
    fuente_subtitulo = Font(name="Calibri", size=10, color=SLATE)
    fuente_marca = Font(name="Calibri", size=9, italic=True, color=CYAN)
    fuente_dato = Font(name="Calibri", size=9, color=SLATE)
    fuente_seccion = Font(name="Calibri", size=11, bold=True, color=NAVY)
    fuente_etiqueta = Font(name="Calibri", size=9.5, bold=True, color=SLATE)
    fuente_celda = Font(name="Calibri", size=10, color=TEXTO)
    fuente_header = Font(name="Calibri", size=10, bold=True, color=BLANCO)
    fuente_total = Font(name="Calibri", size=12, bold=True, color=NAVY)

    relleno_header = PatternFill("solid", fgColor=NAVY)
    relleno_cebra = PatternFill("solid", fgColor=GRIS_CLARO)

    fila = 1

    def escribir(col, texto, fuente=None, alineacion=None, relleno=None):
        celda = ws.cell(row=fila, column=col, value=texto)
        if fuente:
            celda.font = fuente
        if alineacion:
            celda.alignment = alineacion
        if relleno:
            celda.fill = relleno
        return celda

    # ---------- Membrete ----------
    escribir(1, NOMBRE_EMPRESA, fuente_titulo)
    fila += 1
    escribir(1, "Tecnología · Innovación · Calidad", fuente_marca)
    fila += 1
    escribir(1, f"NIT {NIT_EMPRESA}  ·  {REGIMEN_EMPRESA}", fuente_dato)
    fila += 1
    escribir(1, DIRECCION_EMPRESA, fuente_dato)
    fila += 1
    escribir(1, f"{TELEFONO_EMPRESA}  ·  {CORREO_EMPRESA}", fuente_dato)
    fila += 2

    # ---------- Título de la factura ----------
    escribir(1, "Factura de venta electrónica", fuente_titulo)
    fila += 1
    escribir(
        1,
        f"N.º {numero_factura} · Pedido #{numero_pedido} · {formatear_fecha(pedido.get('fecha'))}",
        fuente_subtitulo,
    )
    fila += 2

    # ---------- Datos del cliente ----------
    etiqueta_estado = "Estado de la factura" if es_factura else "Estado del pedido"
    datos_cliente = [
        ("Cliente", nombre_cliente),
        ("Documento", documento_cliente),
        ("Correo", usuario.get("correo") or "—"),
        ("Teléfono", usuario.get("telefono") or "—"),
        ("Dirección", usuario.get("direccion") or "—"),
        (etiqueta_estado, pedido.get("estado") or "—"),
    ]
    for etiqueta, valor in datos_cliente:
        escribir(1, etiqueta, fuente_etiqueta)
        escribir(2, valor, fuente_celda)
        fila += 1
    fila += 1

    # ---------- Tabla de productos y servicios ----------
    escribir(1, "Productos y servicios", fuente_seccion)
    fila += 1

    encabezados = ["Tipo", "Ítem", "Cantidad", "Precio unitario", "Subtotal"]
    for col, encabezado in enumerate(encabezados, start=1):
        alineacion = Alignment(horizontal="left" if col <= 2 else "right")
        escribir(col, encabezado, fuente_header, alineacion, relleno_header)
    fila += 1

    items = pedido.get("productos", []) or []
    for idx, item in enumerate(items):
        relleno = relleno_cebra if idx % 2 == 1 else None
        es_servicio = item.get("tipo") == "servicio"
        nombre_defecto = "Servicio eliminado" if es_servicio else "Producto eliminado"
        escribir(1, "Servicio" if es_servicio else "Producto", fuente_celda, relleno=relleno)
        escribir(2, item.get("nombre") or nombre_defecto, fuente_celda, relleno=relleno)
        escribir(3, item.get("cantidad", "—"), fuente_celda, Alignment(horizontal="right"), relleno)
        escribir(4, formatear_moneda(item.get("precio")), fuente_celda, Alignment(horizontal="right"), relleno)
        escribir(5, formatear_moneda(item.get("subtotal")), fuente_celda, Alignment(horizontal="right"), relleno)
        fila += 1

    if not items:
        escribir(1, "Este pedido no tiene productos ni servicios.", fuente_dato)
        fila += 1

    fila += 1

    # ---------- Totales (Subtotal + Descuento + IVA + Total, mismo criterio que el PDF) ----------
    for etiqueta, valor in lineas_totales_pedido(pedido):
        es_total = etiqueta == "TOTAL"
        escribir(4, etiqueta, fuente_total if es_total else fuente_etiqueta, Alignment(horizontal="right"))
        escribir(5, valor, fuente_total if es_total else fuente_celda, Alignment(horizontal="right"))
        fila += 1

    fila += 2
    generado = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    escribir(1, f"{NOMBRE_EMPRESA} · Comprobante #{numero_pedido} · Generado el {generado}", fuente_dato)

    # ---------- Anchos de columna ----------
    for col, ancho in enumerate([12, 34, 12, 18, 16], start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
